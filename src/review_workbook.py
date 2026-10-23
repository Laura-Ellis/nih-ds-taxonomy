"""Format the hand-coding workbook: instructions sheet, dropdowns, widths, frozen header."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

INPUT_COLS = ["my_ds_related", "my_ai_related", "my_categories", "missed_terms",
              "false_positive_terms", "notes"]

FIELDS = [
    # field, who fills it, definition, allowed values, example
    ("appl_id", "Given", "NIH application ID; unique per project and fiscal year", "", "10012345"),
    ("project_num", "Given", "Full NIH project number", "", "5R01CA123456-03"),
    ("fiscal_year", "Given", "NIH fiscal year (October to September)", "", "2025"),
    ("ic", "Given", "Administering NIH Institute or Center", "", "NCI"),
    ("activity_code", "Given", "Funding mechanism", "", "R01"),
    ("stratum", "Given", "How the project was sampled", "ai_search | general | training", "general"),
    ("project_title, abstract_text", "Given", "Text the labels are based on", "", ""),
    ("rcdc_ml_ai", "Given", "NIH's RCDC category 'Machine Learning and Artificial Intelligence' includes this project", "TRUE | FALSE", "TRUE"),
    ("data_science_related", "Given", "Keyword label: any data science methods, software, infrastructure, or data sharing category", "TRUE | FALSE", "TRUE"),
    ("ai_related", "Given", "Keyword label: any method in the artificial intelligence group", "TRUE | FALSE", "FALSE"),
    ("categories", "Given", "Categories the keyword labeler assigned", "category IDs separated by ;", "machine_learning; clinical_care"),
    ("matched_terms", "Given", "Terms that triggered each category", "category: terms", "machine_learning: random forest"),
    ("my_ds_related", "You", "Does the project use or build data science as part of its own work? Y = data science is in the aims or approach (methods, software, infrastructure, or data sharing). M = mentioned only (background, a core's capabilities, a future possibility). N = none, or only routine statistics such as t-tests and regression.", "Y | M | N", "Y"),
    ("my_ai_related", "You", "Does the project use, develop, or study AI (machine learning, deep learning, NLP, computer vision, generative AI, other AI)? Same Y / M / N rule as my_ds_related. If my_ds_related is N, this is N.", "Y | M | N", "M"),
    ("my_categories", "You", "Every category that applies, from the Category IDs sheet, including application areas. Mark a category only if it is part of the project's own work, not a mention. Leave blank if none.", "category IDs separated by ;", "deep_learning; medical_imaging; research_software"),
    ("missed_terms", "You", "Words or phrases in the abstract that signal a category the keyword labels missed. Use these to add terms.", "category_id: term; category_id: term", "nlp: free-text radiology reports"),
    ("false_positive_terms", "You", "Matched terms that produced a wrong label. Copy from matched_terms. Use these to drop or restrict terms.", "category_id: term", "computer_vision: super-resolution"),
    ("notes", "You", "Anything else: ambiguous cases, a category that seems to be missing, or why you chose M", "free text", "AI only in the core's training plan"),
]

HOW_TO = [
    "How to code",
    "1. Read the title and abstract before looking at the keyword columns, so they do not anchor your judgment.",
    "2. Fill my_ds_related and my_ai_related first (dropdowns), then my_categories.",
    "3. Then compare with categories and matched_terms, and fill missed_terms and false_positive_terms.",
    "4. Code all 150 rows; about 5 minutes each. Ambiguous? Pick the closest answer and explain in notes.",
    "",
    "What the results are used for",
    "Agreement between your coding and the keyword labels (and RCDC) is reported as Cohen's kappa:",
    "κ = (observed agreement − chance agreement) / (1 − chance agreement). 1 = perfect, 0 = chance.",
    "For kappa, M counts as N (not substantive). The M answers also show how often keywords catch passing mentions.",
]


def format_review_workbook(path: Path) -> None:
    wb = load_workbook(path)
    ws = wb["Review"]
    header = {c.value: c.column_letter for c in ws[1]}
    n = ws.max_row

    yes_fill = PatternFill("solid", fgColor="FFF2CC")
    for col in INPUT_COLS:
        letter = header[col]
        ws[f"{letter}1"].fill = PatternFill("solid", fgColor="E1B781")
        for r in range(2, n + 1):
            ws[f"{letter}{r}"].fill = yes_fill
    dv = DataValidation(type="list", formula1='"Y,M,N"', allow_blank=True,
                        error="Use Y, M, or N", showErrorMessage=True)
    ws.add_data_validation(dv)
    for col in ["my_ds_related", "my_ai_related"]:
        dv.add(f"{header[col]}2:{header[col]}{n}")

    widths = {"project_title": 40, "abstract_text": 90, "categories": 30, "matched_terms": 40,
              "my_categories": 30, "missed_terms": 30, "false_positive_terms": 30, "notes": 35}
    for col, letter in header.items():
        ws.column_dimensions[letter].width = widths.get(col, 14)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "B2"

    ins = wb.create_sheet("Instructions", 0)
    ins.append(["Field", "Who fills it", "Definition", "Allowed values", "Example"])
    for f in FIELDS:
        ins.append(list(f))
    ins.append([])
    for line in HOW_TO:
        ins.append([line])
    for c in ins[1]:
        c.font = Font(bold=True)
    for r in range(2, len(FIELDS) + 2):
        if ins[f"B{r}"].value == "You":
            for col in "ABCDE":
                ins[f"{col}{r}"].fill = yes_fill
    for col, w in zip("ABCDE", [26, 12, 80, 30, 40]):
        ins.column_dimensions[col].width = w
    for row in ins.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in (len(FIELDS) + 3, len(FIELDS) + 9):
        ins[f"A{r}"].font = Font(bold=True)
    ins.freeze_panes = "A2"
    wb.active = 0
    wb.save(path)
